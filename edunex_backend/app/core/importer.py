import csv
import io

MAX_ROWS = 5000

ALIASES = {
    "register_number": "register_number", "register_no": "register_number",
    "reg_no": "register_number", "regno": "register_number",
    "registerno": "register_number", "reg_number": "register_number",
    "name": "name", "student_name": "name",
    "email": "email", "college_email": "email", "email_id": "email",
    "class": "class", "class_name": "class", "section": "class",
    "subject_code": "subject_code", "subject": "subject_code",
    "course_code": "subject_code", "code": "subject_code",
}


def _norm(h) -> str:
    return str(h or "").strip().lower().replace(" ", "_").replace("-", "_").replace(".", "")


def _clean(raw_rows: list[dict]) -> list[dict]:
    out = []
    for r in raw_rows:
        row = {}
        for k, v in r.items():
            key = ALIASES.get(_norm(k))
            if key and not isinstance(v, list):
                row[key] = ("" if v is None else str(v)).strip()
        if any(row.values()):
            out.append(row)
    if len(out) > MAX_ROWS:
        raise ValueError(f"Too many rows (max {MAX_ROWS}).")
    return out


def rows_from_text(text: str) -> list[dict]:
    text = text.strip("\ufeff\n\r ")
    if not text:
        raise ValueError("The pasted text is empty.")
    first = text.splitlines()[0]
    delim = max([",", "\t", ";"], key=first.count)   # works for CSV and for text pasted from Excel
    return _clean(list(csv.DictReader(io.StringIO(text), delimiter=delim)))


def rows_from_file(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return rows_from_text(data.decode("utf-8-sig", errors="replace"))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("The spreadsheet is empty.")
        headers = [str(h or "") for h in rows[0]]
        return _clean([dict(zip(headers, r)) for r in rows[1:]])
    raise ValueError("Upload a .csv or .xlsx file.")